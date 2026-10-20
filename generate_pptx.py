import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    # Color Palette: Minimalist White Theme
    COLOR_BG = RGBColor(255, 255, 255)
    COLOR_TEXT = RGBColor(15, 23, 42)      # Slate 900
    COLOR_MUTED = RGBColor(71, 85, 105)    # Slate 600
    COLOR_PRIMARY = RGBColor(30, 58, 138)  # Deep Navy Blue
    COLOR_BLUE = RGBColor(37, 99, 235)     # Primary Blue
    COLOR_EMERALD = RGBColor(5, 150, 105)  # Emerald Green
    COLOR_CARD = RGBColor(248, 250, 252)   # Slate 50
    COLOR_BORDER = RGBColor(226, 232, 240) # Slate 200

    def add_base_slide(cat_text, title_text, slide_num, notes_text=""):
        slide = prs.slides.add_slide(blank_layout)

        # White Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()

        # Header Box (Category & Title)
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)

        # Category / Chapter indicator
        p_cat = tf.paragraphs[0]
        p_cat.text = cat_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_PRIMARY
        p_cat.space_after = Pt(4)

        # Slide Title
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT

        # Footer Box (Divider + Attribution + Slide Number)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_BORDER
        line.line.fill.background()

        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.4))
        ftf = footer_box.text_frame
        ftf.margin_top = Inches(0)
        ftf.margin_left = Inches(0)
        fp = ftf.paragraphs[0]
        fp.text = f"Metode Numerik — Imron dkk. (2022) | Slide {slide_num} dari 21"
        fp.font.size = Pt(10)
        fp.font.color.rgb = COLOR_MUTED

        # Speaker notes removed as requested

        return slide

    def add_card(slide, left, top, width, height, title, items, badge=""):
        # Card background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        # Content box
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), height - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)

        # Card Title
        p_title = tf.paragraphs[0]
        p_title.text = (f"[{badge}] " if badge else "") + title
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_after = Pt(8)

        # Items
        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT
            p.space_after = Pt(6)

    def add_math_box(slide, left, top, width, height, label, math_text, subtext=""):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = COLOR_BLUE
        box.line.width = Pt(2)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.15), width - Inches(0.5), height - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)

        p1 = tf.paragraphs[0]
        p1.text = label.upper()
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_BLUE
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = math_text
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT
        p2.alignment = PP_ALIGN.CENTER
        p2.space_after = Pt(4)

        if subtext:
            p3 = tf.add_paragraph()
            p3.text = subtext
            p3.font.size = Pt(11)
            p3.font.color.rgb = COLOR_MUTED
            p3.alignment = PP_ALIGN.CENTER

    def add_table_data(slide, left, top, width, height, headers, rows):
        num_rows = len(rows) + 1
        num_cols = len(headers)
        table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
        table = table_shape.table

        # Format Header
        for c_idx, h in enumerate(headers):
            cell = table.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD
            p = cell.text_frame.paragraphs[0]
            p.text = str(h)
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_PRIMARY
            p.alignment = PP_ALIGN.CENTER

        # Format Rows
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx + 1, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_BG if r_idx % 2 == 0 else COLOR_CARD
                p = cell.text_frame.paragraphs[0]
                p.text = str(val)
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_TEXT
                p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 1: COVER
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG
    bg1.line.fill.background()

    # Left Accent Bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.12), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(9.5), Inches(2.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "METODE NUMERIK"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "DIFERENSIASI NUMERIK"
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT
    p2.space_after = Pt(6)

    p3 = tf1.add_paragraph()
    p3.text = "Pendekatan Nilai Turunan Fungsi melalui Polinomial Deret & Ekstrapolasi Richardson"
    p3.font.size = Pt(16)
    p3.font.color.rgb = COLOR_MUTED

    # Meta box
    meta_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.1), Inches(9.2), Inches(2.5))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = COLOR_CARD
    meta_box.line.color.rgb = COLOR_BORDER
    mtf = meta_box.text_frame
    mtf.word_wrap = True
    
    mp0 = mtf.paragraphs[0]
    mp0.text = "ANGGOTA KELOMPOK:"
    mp0.font.size = Pt(11)
    mp0.font.bold = True
    mp0.font.color.rgb = COLOR_PRIMARY
    mp0.space_after = Pt(3)
    
    mp1 = mtf.add_paragraph()
    mp1.text = "1. Farel Luther Siagian (4253250041)       2. Fatih Taqiyyuddin (4253250057)\n3. Siti Zulayka Maulida Amali (4253250001)   4. Keegan Gibran Jehian (4253250029)"
    mp1.font.size = Pt(12)
    mp1.font.color.rgb = COLOR_TEXT
    mp1.space_after = Pt(8)
    
    mp2 = mtf.add_paragraph()
    mp2.text = "DOSEN PENGAMPU:"
    mp2.font.size = Pt(11)
    mp2.font.bold = True
    mp2.font.color.rgb = COLOR_PRIMARY
    mp2.space_after = Pt(2)
    
    mp3 = mtf.add_paragraph()
    mp3.text = "Debi Yandra Niska, S.Kom., M.Kom.    |    Tahun: 2026"
    mp3.font.size = Pt(12)
    mp3.font.bold = True
    mp3.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # SLIDE 2: PENGERTIAN
    # =========================================================================
    s2 = add_base_slide("Pengantar", "Pengertian Diferensiasi Numerik", 2,
        "Turunan menyatakan laju perubahan. Kadang fungsi sulit diturunkan, atau kita hanya punya tabel data. Contoh di awal pembahasan adalah fungsi jarak bola jatuh dengan ln cosh; kecepatannya sulit diturunkan secara analitik. Karena itu kita pakai pendekatan numerik.")
    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(3.2), "Pengertian", [
        "Memperkirakan nilai turunan fungsi f'(x) dari titik-titik diskret.",
        "Dipakai saat turunan analitis sulit diperoleh secara eksak.",
        "Metode utama saat data berupa tabel pengukuran."
    ], "A")
    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(3.2), "Tujuan", [
        "Memperkirakan laju perubahan sesaat fungsi.",
        "Menghitung turunan dari data tabel diskret.",
        "Membantu penyelesaian persamaan diferensial."
    ], "B")
    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.733), Inches(3.2), "Analitis vs Numerik", [
        "Analitis: Hasil diferensiasi matematis formal (eksak).",
        "Numerik: Hasil hampiran (pendekatan) bernilai numerik.",
        "Memerlukan pemahaman atas galat hampiran."
    ], "C")
    add_math_box(s2, Inches(0.8), Inches(5.2), Inches(11.733), Inches(1.4),
        "Contoh Motivasi Fisika (hlm. 185)",
        "Jarak jatuh:  z(t) = (m / cd) · ln cosh( √(g·cd / m) · t )",
        "Kecepatan v(t) = z'(t) sulit diturunkan analitis pada model aerodinamika non-linier kompleks.")

    # =========================================================================
    # SLIDE 3: PENERAPAN DI DUNIA NYATA
    # =========================================================================
    s3 = add_base_slide("Pengantar Tambahan", "Penerapan di Dunia Nyata", 3,
        "Slide ini penjelasan tambahan agar audiens tahu untuk apa rumus diferensiasi numerik dipakai. Datanya sengaja dibuat sederhana dari s(t) = t^3 supaya kita tahu jawaban benarnya (12 m/s).")
    add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.7), "Contoh Aplikasi Rekayasa", [
        "Kecepatan: v(t) = ds/dt dari sensor posisi GPS / telemetri.",
        "Percepatan: a(t) = dv/dt dari data kecepatan kendaraan.",
        "Perpindahan Panas: Laju perubahan suhu dT/dx (Hukum Fourier).",
        "Data Eksperimen: Menghitung laju reaksi kimia dari sampel waktu.",
        "Data nyata tidak memiliki rumus fungsi kontinu eksak!"
    ], "Praktis")
    add_card(s3, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.7), "Studi Kasus: Posisi s(t) -> v(2)", [
        "Data t = 1, 2, 3 s -> s = 1, 8, 27 m (h = 1 s). Dicari v(2):",
        "Maju: (27 - 8)/1 = 19 m/s (Galat = 7)",
        "Mundur: (8 - 1)/1 = 7 m/s (Galat = 5)",
        "Tengah: (27 - 1)/2 = 13 m/s (Galat = 1)",
        "Eksak s(t)=t³ -> s'(2) = 3(2)² = 12 m/s.",
        "Kesimpulan: Selisih tengah O(h²) jauh lebih teliti daripada O(h)!"
    ], "Kasus")

    # =========================================================================
    # SLIDE 4: DASAR RUMUS: PENDEKATAN DERET
    # =========================================================================
    s4 = add_base_slide("Pendekatan Deret", "Dasar Rumus: Pendekatan Deret", 4,
        "Buku menurunkan rumus lewat polinomial Lagrange dan suku galatnya, bukan langsung lewat Taylor. Intinya: fungsi didekati dengan polinomial yang melalui titik data, lalu diturunkan.")
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.7), "Penurunan via Polinomial Lagrange", [
        "Interpolasi pada x₀ dan x₁ = x₀ + h:",
        "f(x) = P₀,₁(x) + [(x - x₀)(x - x₁)/2!] f''(ξ(x))",
        "P₀,₁(x) = [f(x₀)(x - x₀ - h)/(-h)] + [f(x₀+h)(x - x₀)/h]",
        "Turunkan terhadap x, lalu evaluasi di x = x₀.",
        "Suku koefisien D_x(f''(ξ(x))) bernilai 0 di x = x₀ sehingga galat dapat dirumuskan secara tegas."
    ], "Penurunan")
    add_math_box(s4, Inches(6.8), Inches(1.8), Inches(5.733), Inches(2.2),
        "Hasil Penurunan (Persamaan 6.1 Buku)",
        "f'(x₀) = [f(x₀ + h) - f(x₀)] / h  -  (h / 2) f''(ξ)",
        "Suku -(h/2)f''(ξ) adalah galat pemotongan (truncation error)")
    add_card(s4, Inches(6.8), Inches(4.3), Inches(5.733), Inches(2.2), "Sudut Pandang Deret Taylor", [
        "f(x+h) = f(x) + h f'(x) + (h²/2) f''(x) + ...",
        "f'(x) = [f(x+h) - f(x)] / h - O(h)",
        "Alur: Titik Data -> Polinomial -> Rumus -> Suku Galat"
    ], "Setara")

    # =========================================================================
    # SLIDE 5: SELISIH MAJU DAN MUNDUR
    # =========================================================================
    s5 = add_base_slide("Pendekatan Deret", "Metode Selisih Maju dan Mundur", 5,
        "Dua rumus ini adalah yang paling sederhana. Maju memakai titik di depan, mundur memakai titik di belakang. Keduanya hanya butuh dua nilai fungsi dan galatnya berorde h.")
    add_math_box(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.6),
        "Selisih Maju (Forward Difference)",
        "f'(x) ≈ [f(x + h) - f(x)] / h",
        "Titik di depan | Cocok untuk ujung kiri data | Galat O(h)")
    add_math_box(s5, Inches(6.8), Inches(1.8), Inches(5.733), Inches(2.6),
        "Selisih Mundur (Backward Difference)",
        "f'(x) ≈ [f(x) - f(x - h)] / h",
        "Titik di belakang | Cocok untuk ujung kanan data | Galat O(h)")
    add_card(s5, Inches(0.8), Inches(4.7), Inches(11.733), Inches(1.8), "Karakteristik & Kesimpulan Praktis", [
        "Sederhana: Masing-masing hanya membutuhkan 2 nilai fungsi.",
        "Keterbatasan: Orde galat O(h) memerlukan nilai h yang sangat kecil untuk presisi tinggi.",
        "Buku teks: Beda mundur secara matematis adalah Persamaan 6.1 dengan nilai h negatif."
    ], "Evaluasi")

    # =========================================================================
    # SLIDE 6: SELISIH TENGAH
    # =========================================================================
    s6 = add_base_slide("Pendekatan Deret", "Metode Selisih Tengah (Central Difference)", 6,
        "Selisih tengah adalah rumus tiga titik tengah di buku (6.5), tanpa memakai f(x₀) itu sendiri. Karena simetris, suku galat orde pertama saling meniadakan, sehingga galatnya berorde h².")
    add_math_box(s6, Inches(0.8), Inches(1.8), Inches(11.733), Inches(2.5),
        "Rumus Tiga Titik Tengah (Persamaan 6.5)",
        "f'(x₀) ≈ [f(x₀ + h) - f(x₀ - h)] / (2h)",
        "Suku galat pemotongan:  -(h² / 6) f'''(ξ₁) ,  ξ₁ ∈ (x₀ - h, x₀ + h)  =>  O(h²)")
    add_card(s6, Inches(0.8), Inches(4.6), Inches(5.6), Inches(1.9), "Kelebihan", [
        "Simetris: Suku ganjil Taylor saling meniadakan.",
        "Orde O(h²): Jika h dibagi 2, galat turun 4 kali lipat!",
        "Nilai f(x₀) tidak perlu dihitung dalam rumus."
    ], "Kelebihan")
    add_card(s6, Inches(6.8), Inches(4.6), Inches(5.733), Inches(1.9), "Keterbatasan", [
        "Membutuhkan nilai fungsi di kedua sisi titik (kiri & kanan).",
        "Tidak dapat digunakan pada batas paling ujung data tabel.",
        "Memerlukan rumus titik ujung (slide 7)."
    ], "Keterbatasan")

    # =========================================================================
    # SLIDE 7: RUMUS TIGA TITIK DAN LIMA TITIK
    # =========================================================================
    s7 = add_base_slide("Pendekatan Deret", "Rumus Tiga Titik dan Lima Titik", 7,
        "Rumus tiga titik ujung dipakai saat titik ada di ujung data. Rumus lima titik tengah memakai koefisien 1, -8, 8, -1 dan penyebut 12h; galatnya berorde h⁴.")
    add_math_box(s7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(2.3),
        "3 Titik Ujung (6.4)",
        "[-3f₀ + 4f₁ - f₂] / 2h",
        "Galat: (h²/3) f''' => O(h²)")
    add_math_box(s7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(2.3),
        "5 Titik Tengah (6.7)",
        "[f₋₂ - 8f₋₁ + 8f₁ - f₂] / 12h",
        "Galat: (h⁴/30) f⁽⁵⁾ => O(h⁴)")
    add_math_box(s7, Inches(8.8), Inches(1.8), Inches(3.733), Inches(2.3),
        "5 Titik Ujung (6.6)",
        "[-25f₀ + 48f₁ - 36f₂ + 16f₃ - 3f₄] / 12h",
        "Galat: (h⁴/5) f⁽⁵⁾ => O(h⁴)")

    add_table_data(s7, Inches(0.8), Inches(4.4), Inches(11.733), Inches(2.1),
        ["Aspek", "Tiga Titik Ujung (6.4)", "Lima Titik Tengah (6.7)", "Lima Titik Ujung (6.6)"],
        [
            ["Jumlah Titik", "3 titik", "5 titik", "5 titik"],
            ["Posisi", "Ujung batas interval", "Titik dalam (interior)", "Ujung batas interval"],
            ["Orde Galat", "O(h²)", "O(h⁴) ⭐", "O(h⁴) ⭐"]
        ])

    # =========================================================================
    # SLIDE 8: CONTOH 6.1.1 f(x) = ln x
    # =========================================================================
    s8 = add_base_slide("Contoh 6.1.1", "Contoh: Selisih Maju untuk f(x) = ln x", 8,
        "Substitusikan x₀ = 1,8 dan h = 0,1 ke rumus beda maju: ln 1,9 dikurangi ln 1,8, dibagi 0,1, hasilnya 0,5406722. Bandingkan dengan eksak 1/1,8 = 0,5555556.")
    add_math_box(s8, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.8),
        "Perhitungan Kasus x₀ = 1,8 ; h = 0,1",
        "f'(1,8) ≈ [ln(1,9) - ln(1,8)] / 0,1 = (0,64185389 - 0,58778667) / 0,1 = 0,5406722",
        "Nilai Eksak: f'(x) = 1/x => f'(1,8) = 1/1,8 ≈ 0,5555556")
    add_table_data(s8, Inches(0.8), Inches(3.8), Inches(11.733), Inches(2.2),
        ["Langkah h", "Nilai Pendekatan", "Galat Aktual |Eksak - Pend.|", "Batas Atas Galat |h| / [2(1,8)²]"],
        [
            ["0,10", "0,5406722", "0,0148834", "0,0154321"],
            ["0,05", "0,5479795", "0,0075761", "0,0077160"],
            ["0,01", "0,5540180", "0,0015376 ⭐", "0,0015432"]
        ])
    add_card(s8, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.65), "Keterangan", [
        "Makin kecil h, pendekatan makin dekat ke eksak. Batas atas teoritis terbukti selalu lebih besar dari galat aktual."
    ])

    # =========================================================================
    # SLIDE 9: ANALISIS GALAT
    # =========================================================================
    s9 = add_base_slide("Analisis Galat", "Galat atau Kesalahan Pendekatan", 9,
        "Setiap rumus yang kita pakai hanya pendekatan. Galat pemotongan muncul karena suku dibuang. Orde O(h), O(h²), O(h⁴) menunjukkan kecepatan galat mengecil.")
    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.7), "Konsep Dasar Galat", [
        "Turunan numerik selalu berupa hampiran / pendekatan.",
        "Galat Pemotongan (Truncation Error): Akibat memutus deret tak hingga Taylor.",
        "Galat Bertanda = Eksak - Pendekatan.",
        "Galat Absolut = |Eksak - Pendekatan|.",
        "Pengaruh h: h makin kecil -> galat pemotongan umumnya turun."
    ], "Teori")
    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.7), "Perbandingan Orde Galat O(h^p)", [
        "Selisih Maju  : O(h)    [Linear, penurunan lambat]",
        "Selisih Mundur: O(h)    [Linear, penurunan lambat]",
        "Selisih Tengah: O(h²)   [Kuadratik, h/2 -> galat / 4]",
        "3 Titik Ujung : O(h²)   [Kuadratik, h/2 -> galat / 4]",
        "5 Titik Tengah: O(h⁴)   [Kuartik, h/2 -> galat / 16! ⭐]",
        "5 Titik Ujung : O(h⁴)   [Kuartik, h/2 -> galat / 16! ⭐]",
        "Notasi Big-O adalah perilaku asimtotik saat h cukup kecil."
    ], "Orde")

    # =========================================================================
    # SLIDE 10: PENGARUH h & DILEMA PEMBULATAN
    # =========================================================================
    s10 = add_base_slide("Analisis Galat Tambahan", "Pengaruh Nilai h dan Sumber Galat", 10,
        "Ada dua sumber galat berlawanan: pemotongan (mengecil saat h turun) vs pembulatan/noise (membesar saat h turun). Tabel dihitung sendiri dengan e^x di x=1.")
    add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.7), "Trade-off Komputasi", [
        "Galat Pemotongan ~ h^p: Turun saat h mengecil.",
        "Galat Pembulatan / Noise ~ δ/h: Naik saat h mengecil!",
        "Total Galat (Tengah): E(h) ≤ (M h² / 6) + (δ / h)",
        "Langkah Optimum: h* ≈ ³√(3δ / M)",
        "Peringatan: Nilai h yang terlalu kecil TIDAK SELALU lebih akurat!"
    ], "Model")
    add_table_data(s10, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.7),
        ["Langkah h", "Galat Maju", "Galat Tengah"],
        [
            ["10⁻¹", "1,41 × 10⁻¹", "4,53 × 10⁻³"],
            ["10⁻⁴", "1,36 × 10⁻⁴", "4,53 × 10⁻⁹"],
            ["10⁻⁶", "1,36 × 10⁻⁶", "1,63 × 10⁻¹⁰ (Optimum Tengah) ⭐"],
            ["10⁻⁸", "6,60 × 10⁻⁹ (Optimum Maju) ⭐", "6,60 × 10⁻⁹"],
            ["10⁻¹²", "4,32 × 10⁻⁴ (Naik kembali)", "2,10 × 10⁻⁴ (Naik kembali)"],
            ["10⁻¹⁴", "9,34 × 10⁻³ (Naik kembali)", "9,34 × 10⁻³ (Naik kembali)"]
        ])

    # =========================================================================
    # SLIDE 11: GALAT ABSOLUT, RELATIF, PERSENTASE
    # =========================================================================
    s11 = add_base_slide("Analisis Galat Tambahan", "Galat Absolut, Relatif, dan Persentase", 11,
        "Slide ini melengkapi slide galat: Ea memberi satuan turunan, Er membandingkan dengan besar turunan, E_% adalah Er x 100. Jika turunan nol, Er tidak terdefinisi.")
    add_math_box(s11, Inches(0.8), Inches(1.8), Inches(3.6), Inches(2.2),
        "Galat Absolut (Ea)",
        "Ea = |f'(x) - Dh|",
        "Satuan sama dengan turunan f'")
    add_math_box(s11, Inches(4.8), Inches(1.8), Inches(3.6), Inches(2.2),
        "Galat Relatif (Er)",
        "Er = |f'(x) - Dh| / |f'(x)|",
        "Ukuran proporsional")
    add_math_box(s11, Inches(8.8), Inches(1.8), Inches(3.733), Inches(2.2),
        "Galat Persentase (E%)",
        "E% = Er × 100%",
        "Format persentase praktis")

    add_card(s11, Inches(0.8), Inches(4.3), Inches(11.733), Inches(2.3), "Contoh Aplikasi: Kasus ln x di x = 1,8 (h = 0,1)", [
        "Dh = 0,5406722  dan  f'(1,8) = 0,5555556",
        "Ea = |0,5555556 - 0,5406722| = 0,0148834",
        "Er = 0,0148834 / 0,5555556 ≈ 0,0267901",
        "E% ≈ 2,679%",
        "Perhatian Khusus: Jika f'(x) = 0, Er tidak terdefinisi (pembagian nol). Gunakan Ea!"
    ], "Aplikasi")

    # =========================================================================
    # SLIDE 12: EKSTRAPOLASI RICHARDSON
    # =========================================================================
    s12 = add_base_slide("Ekstrapolasi Richardson", "Metode Ekstrapolasi Richardson", 12,
        "Ide Richardson: dua hasil dengan h berbeda digabung agar suku galat utama hilang. Kunci: periksa dulu bentuk galat awalnya, baru pilih rumus.")
    add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.7), "Prinsip Kerja Richardson", [
        "Menggabungkan hasil langkah awal N₁(h) dan langkah separuhnya N₁(h/2).",
        "Menghilangkan suku galat dominan pertama.",
        "Mendongkrak ketelitian secara drastis tanpa perlu rumus baru.",
        "Rumus Rekursif Umum (Galat Genap):",
        "Nj(h) = N_{j-1}(h/2) + [1/(4^{j-1} - 1)] [N_{j-1}(h/2) - N_{j-1}(h)]",
        "Penting: Pilih rumus sesuai bentuk ekspansi galat!"
    ], "Konsep")
    add_math_box(s12, Inches(6.8), Inches(1.8), Inches(5.733), Inches(2.2),
        "Galat Berpangkat Genap (Selisih Tengah)",
        "N₂(h) = [4 N₁(h/2) - N₁(h)] / 3",
        "Menghilangkan suku h² => Menjadi berorde O(h⁴) ⭐")
    add_math_box(s12, Inches(6.8), Inches(4.3), Inches(5.733), Inches(2.2),
        "Galat Berpangkat Linear (Selisih Maju/Mundur)",
        "N₂(h) = 2 N₁(h/2) - N₁(h)",
        "Menghilangkan suku h => Menjadi berorde O(h²)")

    # =========================================================================
    # SLIDE 13: CONTOH 6.2.1 RICHARDSON f(x) = ln x
    # =========================================================================
    s13 = add_base_slide("Contoh 6.2.1", "Contoh Richardson untuk f(x) = ln x", 13,
        "Ambil dua hasil dari Slide 8: N₁(0,1) dan N₁(0,05). Karena selisih maju berorde h, pakai N₂ = 2N₁(h/2) - N₁(h). Hasilnya jauh lebih dekat ke nilai eksak.")
    add_math_box(s13, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.8),
        "Substitusi Nilai Langkah h = 0,1 dan h = 0,05",
        "N₂(0,1) = 2 N₁(0,05) - N₁(0,1) = 2(0,5479795) - 0,5406722 ≈ 0,5552868",
        "Nilai Eksak: 1/1,8 ≈ 0,5555556")
    add_table_data(s13, Inches(0.8), Inches(3.8), Inches(11.733), Inches(2.7),
        ["Metode", "Hasil Perhitungan", "Galat terhadap Nilai Eksak"],
        [
            ["Selisih Maju (h = 0,1)", "0,5406722", "0,0148834"],
            ["Selisih Maju (h = 0,05)", "0,5479795", "0,0075761"],
            ["Ekstrapolasi Richardson", "0,5552868", "0,0002688 (28× Lebih Teliti!) ⭐"],
            ["Nilai Eksak", "0,5555556", "— (Referensi)"]
        ])

    # =========================================================================
    # SLIDE 14: CONTOH 6.1.2 DATA TABEL f(x) = x e^x
    # =========================================================================
    s14 = add_base_slide("Contoh 6.1.2", "Penerapan pada Data Tabel: f(x) = x e^x", 14,
        "Data hanya tabel lima titik. Tiga titik ujung: 22,032310. Tiga titik tengah: 22,228790. Lima titik tengah: 22,166999. Nilai eksak 3e² = 22,167168.")
    add_table_data(s14, Inches(0.8), Inches(1.8), Inches(4.5), Inches(4.7),
        ["x", "f(x) Data Tabel"],
        [
            ["1,8", "10,889365"],
            ["1,9", "12,703199"],
            ["2,0 (Titik Evaluasi)", "14,778112"],
            ["2,1", "17,148957"],
            ["2,2", "19,855030"]
        ])
    add_table_data(s14, Inches(5.6), Inches(1.8), Inches(6.933), Inches(3.0),
        ["Rumus (h = 0,1)", "Pendekatan f'(2)", "Galat Absolut"],
        [
            ["Tiga Titik Ujung", "22,032310", "1,35 × 10⁻¹"],
            ["Tiga Titik Tengah", "22,228790", "6,16 × 10⁻²"],
            ["Lima Titik Tengah", "22,166999", "1,69 × 10⁻⁴ ⭐"]
        ])
    add_card(s14, Inches(5.6), Inches(5.0), Inches(6.933), Inches(1.5), "Analisis Hasil", [
        "Eksak: f'(x) = (1+x)e^x => f'(2) = 3e² ≈ 22,167168.",
        "Lima titik tengah O(h⁴) memberikan ketelitian luar biasa pada data tabel!"
    ], "Evaluasi")

    # =========================================================================
    # SLIDE 15: CONTOH 6.1.3 DATA TABEL f'(1,1) & f'(1,3)
    # =========================================================================
    s15 = add_base_slide("Contoh 6.1.3", "Contoh 6.1.3 — Menghitung Turunan dari Data Tabel", 15,
        "Di sini diminta melengkapi kolom turunan dari tabel. Di ujung pakai rumus ujung, di tengah pakai rumus tengah. Lima titik tengah hanya bisa untuk x=1,3.")
    add_table_data(s15, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.7),
        ["Titik x", "Rumus yang Cocok Digunakan (h = 0,1)"],
        [
            ["1,1 (Ujung)", "6.4 Tiga Titik Ujung ; 6.6 Lima Titik Ujung"],
            ["1,2 (Interior)", "6.4 Ujung ; 6.5 Tiga Titik Tengah"],
            ["1,3 (Tengah Penuh)", "6.4 Ujung ; 6.5 Tengah ; 6.7 Lima Titik Tengah ⭐"],
            ["1,4 (Interior)", "6.4 Ujung (h=-0,1) ; 6.5 Tengah"],
            ["1,5 (Ujung)", "6.4 Ujung (h=-0,1) ; 6.6 Lima Titik Ujung"]
        ])
    add_card(s15, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.7), "Contoh Hitungan Tabel", [
        "f'(1,1) [Ujung 6.4, h=0,1]:",
        "[-3(9,025013) + 4(11,02318) - 13,46374] / 0,2 = 17,76971",
        "",
        "f'(1,3) [Tengah 6.5, h=0,1]:",
        "[16,44465 - 11,02318] / 0,2 = 27,10735",
        "",
        "f'(1,3) [Lima Titik Tengah 6.7]:",
        "[9,025013 - 8(11,02318) + 8(16,44465) - 20,08554] / 1,2 = 26,92603 ⭐"
    ], "Hitungan")

    # =========================================================================
    # SLIDE 16: CONTOH 6.2.2 RICHARDSON f(x) = x e^x
    # =========================================================================
    s16 = add_base_slide("Contoh 6.2.2", "Contoh 6.2.2 — Richardson pada f(x) = x e^x", 16,
        "Rumus tengah punya galat berpangkat genap, jadi dipakai pembagi 3 lalu 15. N₃(0,2) sama dengan nilai eksak sampai 6 angka desimal.")
    add_table_data(s16, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.0),
        ["Langkah / Tingkat", "Perhitungan", "Hasil", "Keterangan"],
        [
            ["N₁(0,2)", "2,5(19,855030 - 10,889365)", "22,414160", "Selisih Tengah h=0,2"],
            ["N₁(0,1)", "5(17,148957 - 12,703199)", "22,228786", "Selisih Tengah h=0,1"],
            ["N₂(0,2)", "22,228786 + (22,228786 - 22,414160)/3", "22,166995", "Tingkat 2 (O(h⁴))"],
            ["N₁(0,05)", "10(15,924197 - 13,705941)", "22,182564", "Selisih Tengah h=0,05"],
            ["N₂(0,1)", "22,182564 + (22,182564 - 22,228786)/3", "22,167157", "Tingkat 2 h=0,1"],
            ["N₃(0,2)", "22,167157 + (22,167157 - 22,166995)/15", "22,167168 ⭐", "Identik Eksak!"]
        ])
    add_card(s16, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.75), "Hasil Akhir", [
        "Eksak: 3e² ≈ 22,167168. N₃(0,2) identik sempurna dengan nilai eksak hingga 6 angka desimal!"
    ])

    # =========================================================================
    # SLIDE 17: CONTOH 6.2.3 RICHARDSON DARI TABEL
    # =========================================================================
    s17 = add_base_slide("Contoh 6.2.3", "Contoh 6.2.3 — Richardson dari Data Tabel", 17,
        "Titik yang dihitung f'(1,3) dengan h=0,2 dan 0,1. Level ketiga tidak mungkin dihitung karena data di 1,25 dan 1,35 tidak tersedia.")
    add_card(s17, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.7), "Perhitungan Richardson x₀ = 1,3", [
        "N₁(0,2) = (20,08554 - 9,025013) / 0,4 = 27,651318",
        "N₁(0,1) = (16,44465 - 11,02318) / 0,2 = 27,10735",
        "",
        "Kombinasi Galat Genap N₂(0,2):",
        "N₂(0,2) = 27,10735 + (27,10735 - 27,651318) / 3",
        "N₂(0,2) ≈ 26,926027 ⭐",
        "Hasil ini berorde O(h⁴) dan sama dengan lima titik tengah!"
    ], "Hitungan")
    add_card(s17, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.7), "Keterbatasan Batas Data", [
        "Mengapa N₃ tidak bisa dihitung?",
        "Untuk N₃ dibutuhkan N₁(0,05) yang memerlukan f(1,35) dan f(1,25).",
        "Karena tabel eksperimen hanya memiliki kelipatan 0,1, data tersebut tidak ada.",
        "Rekursi Richardson harus berhenti pada ketersediaan data tabel nyata!"
    ], "Batasan")

    # =========================================================================
    # SLIDE 18: RINGKASAN METODE: PERBANDINGAN
    # =========================================================================
    s18 = add_base_slide("Ringkasan Metode", "Perbandingan Metode Diferensiasi Numerik", 18,
        "Gunakan tabel ini untuk merangkum. Maju/mundur O(h), tengah & 3 titik ujung O(h²), lima titik O(h⁴). Richardson memperbaiki hasil rumus lain.")
    add_table_data(s18, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8),
        ["Metode", "Titik / Data", "Galat", "Kelebihan", "Keterbatasan"],
        [
            ["Selisih Maju", "x, x+h", "O(h)", "Sederhana, 2 nilai", "Kurang teliti"],
            ["Selisih Mundur", "x-h, x", "O(h)", "Cocok di ujung kanan", "Kurang teliti"],
            ["Selisih Tengah", "x-h, x+h", "O(h²)", "Simetris, lebih teliti", "Perlu data 2 sisi"],
            ["3 Titik Ujung", "x₀, x₀+h, x₀+2h", "O(h²)", "Di ujung interval", "Butuh 3 nilai"],
            ["5 Titik Tengah", "x₀±h, x₀±2h", "O(h⁴) ⭐", "Sangat teliti & presisi", "Butuh 5 nilai simetris"],
            ["5 Titik Ujung", "x₀ ... x₀+4h", "O(h⁴) ⭐", "Teliti di ujung satu sisi", "Butuh 5 nilai"],
            ["Richardson", "N₁(h), N₁(h/2)", "Tinggi", "Mendongkrak ketelitian", "Bentuk galat harus sesuai"]
        ])

    # =========================================================================
    # SLIDE 19: MEMILIH METODE & KONDISI
    # =========================================================================
    s19 = add_base_slide("Ringkasan Metode", "Memilih Metode: Rumus & Kondisi Penggunaan", 19,
        "Slide ini fokus pada kapan rumus dipakai. Pesan penutup: jangan menyimpulkan satu metode selalu paling baik. Lima titik memiliki koefisien besar sehingga lebih sensitif terhadap noise.")
    add_table_data(s19, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.3),
        ["Metode", "Rumus Dasar", "Orde", "Kondisi Penggunaan"],
        [
            ["Selisih Maju", "[f(x+h) - f(x)] / h", "O(h)", "Ujung kiri data; hanya ada data ke depan"],
            ["Selisih Mundur", "[f(x) - f(x-h)] / h", "O(h)", "Ujung kanan data; hanya ada data ke belakang"],
            ["Selisih Tengah", "[f(x+h) - f(x-h)] / 2h", "O(h²)", "Titik dalam; pilihan umum paling seimbang"],
            ["3 Titik Ujung", "[-3f₀ + 4f₁ - f₂] / 2h", "O(h²)", "Ujung interval yang membutuhkan akurasi orde 2"],
            ["5 Titik Tengah", "[f₋₂ - 8f₋₁ + 8f₁ - f₂] / 12h", "O(h⁴)", "Titik dalam; data rapat dan fungsi mulus"],
            ["5 Titik Ujung", "[-25f₀ + 48f₁ - 36f₂ + 16f₃ - 3f₄]/12h", "O(h⁴)", "Ujung interval yang membutuhkan akurasi orde 4"],
            ["Richardson", "Eliminasi galat N₁(h) & N₁(h/2)", "Tinggi", "Tersedia multi-langkah & bentuk galat sesuai"]
        ])
    add_card(s19, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.55), "Keterangan", [
        "Tidak ada metode yang selalu terbaik di segala situasi: perhatikan noise data dan kestabilan pembulatan!"
    ])

    # =========================================================================
    # SLIDE 20: KESIMPULAN
    # =========================================================================
    s20 = add_base_slide("Penutup", "Kesimpulan", 20,
        "Rangkum perjalanan dari definisi, deret, rumus 2, 3, 5 titik, galat, lalu Richardson. Tutup dengan kutipan penting.")
    add_card(s20, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.7), "Poin-Poin Kunci", [
        "Diferensiasi numerik memperkirakan turunan dari nilai fungsi diskret.",
        "Pendekatan deret (Lagrange & Taylor) menjadi fondasi matematis utama.",
        "Selisih maju dan mundur: galat pemotongan O(h).",
        "Selisih tengah dan tiga titik ujung: galat O(h²).",
        "Lima titik (tengah dan ujung): ketelitian tinggi O(h⁴).",
        "Ekstrapolasi Richardson meningkatkan ketelitian secara signifikan.",
        "Pilih h dengan cermat mengingat trade-off pemotongan vs pembulatan."
    ], "Ringkasan")
    add_card(s20, Inches(7.2), Inches(1.8), Inches(5.333), Inches(4.7), "Prinsip Utama", [
        "\"Ketelitian diferensiasi numerik tidak hanya ditentukan oleh rumus yang digunakan, tetapi juga oleh ukuran langkah, data yang tersedia, dan sifat fungsi.\"",
        "",
        "Keseimbangan Komputasi:",
        "• Orde tinggi butuh lebih banyak data & sensitif terhadap noise.",
        "• Nilai h terlalu kecil memperbesar galat pembulatan komputer.",
        "• Selalu sesuaikan metode dengan posisi titik data!"
    ], "Intisari")

    # =========================================================================
    # SLIDE 21: DAFTAR PUSTAKA
    # =========================================================================
    s21 = add_base_slide("Referensi", "Daftar Pustaka", 21,
        "Sebutkan sumber utama presentasi: buku Metode Numerik karya Imron, Mardlijah, dan Asfihani (2022), Departemen Matematika ITS.")
    add_card(s21, Inches(0.8), Inches(1.8), Inches(11.733), Inches(2.2), "Buku Teks Utama", [
        "Imron, C., Mardlijah, & Asfihani, T. (2022).",
        "Metode Numerik.",
        "Departemen Matematika, Institut Teknologi Sepuluh Nopember (ITS).",
        "ISBN: 978-623-318-126-6."
    ], "Utama")
    add_card(s21, Inches(0.8), Inches(4.3), Inches(11.733), Inches(2.2), "Referensi Tambahan", [
        "[1] Burden, R. L., & Faires, J. D. (2010). Numerical Analysis, 9th ed. Cengage Learning. §1.2 (Galat), §4.1 (Diferensiasi Numerik), §4.2 (Ekstrapolasi Richardson).",
        "[2] University of Maryland (AMSC 466). Lecture Notes: Numerical Differentiation. Pertukaran galat pemotongan dan pembulatan komputer."
    ], "Pendukung")

    # Save presentation
    output_path = "Diferensiasi_Numerik.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
